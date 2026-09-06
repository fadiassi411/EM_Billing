from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "images"
OUT = ROOT
NAVY = "0B2B3A"; TEAL = "008F83"; AQUA = "DFF3F1"; BLUE = "1976B9"; PALE = "F2F7F8"; GOLD = "D9911B"; RED = "C73B3B"; GREY = "526875"

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = tcPr.find(qn("w:shd"))
    if el is None: el = OxmlElement("w:shd"); tcPr.append(el)
    el.set(qn("w:fill"), color)

def margins(sec):
    sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.68); sec.right_margin=Inches(.68)

def add_page_num(par):
    par.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=par.add_run("PAGE "); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

def picture(run,path,width,alt):
    run.add_picture(str(path),width=Inches(width))
    for prop in run._r.xpath('.//wp:docPr'):
        prop.set('descr',alt); prop.set('title',alt)

def base_doc():
    d=Document(); margins(d.sections[0]);
    styles=d.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9); styles['Normal'].font.color.rgb=RGBColor.from_string(GREY)
    styles['Normal'].paragraph_format.space_after=Pt(5)
    for n,size,color in [('Title',34,NAVY),('Heading 1',22,NAVY),('Heading 2',14,TEAL),('Heading 3',11,NAVY)]:
        s=styles[n]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(8); s.paragraph_format.space_after=Pt(6)
    if 'Caption' in styles:
        styles['Caption'].font.name='Aptos'; styles['Caption'].font.size=Pt(8); styles['Caption'].font.italic=True; styles['Caption'].font.color.rgb=RGBColor.from_string(GREY)
    for sec in d.sections:
        footer=sec.footer.paragraphs[0]; footer.text='WATCH DOG ENERGY MANAGEMENT  •  CUSTOMER USER GUIDE'; footer.runs[0].font.size=Pt(7); footer.runs[0].font.color.rgb=RGBColor.from_string(GREY); add_page_num(footer)
    return d

def label(d,text,color=TEAL):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(3); r=p.add_run(text.upper()); r.bold=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(color)

def title(d,text,subtitle=None):
    d.add_heading(text,0)
    if subtitle:
        p=d.add_paragraph(subtitle); p.style='Subtitle'

def bullet(d,text):
    d.add_paragraph(text,style='List Bullet')

def numbered(d,text):
    d.add_paragraph(text,style='List Number')

def note(d,heading,text,color=AQUA):
    t=d.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade(c,color); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(heading+"  "); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); p.add_run(text)
    d.add_paragraph().paragraph_format.space_after=Pt(0)

def screenshot(d,name,caption,width=7.05):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; picture(p.add_run(),IMG/name,width,caption)
    p=d.add_paragraph(caption,style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def page(d): d.add_page_break()

def simple_table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); r=c.paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    trPr=t.rows[0]._tr.get_or_add_trPr(); header=OxmlElement('w:tblHeader'); header.set(qn('w:val'),'true'); trPr.append(header)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    return t

def diagram(filename,title,boxes,arrows):
    W,H=1500,520; im=Image.new('RGB',(W,H),'white'); dr=ImageDraw.Draw(im)
    try: font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',28); bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',35); small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',22)
    except: font=bold=small=None
    dr.rounded_rectangle((20,20,W-20,H-20),26,fill='#F2F7F8',outline='#C9D9DE',width=3); dr.text((60,50),title,font=bold,fill='#0B2B3A')
    for x,y,w,h,head,body,color in boxes:
        dr.rounded_rectangle((x,y,x+w,y+h),20,fill=color,outline='#0B2B3A',width=3); dr.text((x+20,y+18),head,font=font,fill='#0B2B3A'); dr.multiline_text((x+20,y+63),body,font=small,fill='#526875',spacing=7)
    for x1,y1,x2,y2 in arrows:
        dr.line((x1,y1,x2,y2),fill='#008F83',width=8); dr.polygon([(x2,y2),(x2-20,y2-13),(x2-20,y2+13)],fill='#008F83')
    im.save(IMG/filename)

def make_diagrams():
    # Replace the temporary capture path with customer-safe wording while retaining the real UI.
    backup=Image.open(IMG/'09-backup-restore.png').convert('RGB'); bd=ImageDraw.Draw(backup)
    bd.rectangle((90,55,970,86),fill='white')
    try: bf=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',15)
    except: bf=None
    bd.text((96,60),'Backup folder: [Watch Dog customer data folder]\\Backups',font=bf,fill='#526875')
    backup.save(IMG/'09-backup-restore.png')
    diagram('diagram-system.png','How Watch Dog turns meter readings into invoices',[
      (60,150,250,220,'Meters','Electricity / water\nCumulative readings','#FFFFFF'),(380,150,250,220,'Controller','DDC / PLC or\nprotocol gateway','#FFFFFF'),(700,150,300,220,'Watch Dog PC','Polls once • stores data\ncalculates period usage','#DFF3F1'),(1070,150,360,220,'Customer output','Dashboard • charts\nPDF invoice • optional email','#FFFFFF')],[(310,260,380,260),(630,260,700,260),(1000,260,1070,260)])
    diagram('diagram-lan.png','One Watch Dog system, multiple browser viewers',[
      (70,150,320,220,'Main Watch Dog PC','Backend + database\nMeter polling + port 5080','#DFF3F1'),(560,150,300,220,'Mall LAN / Wi-Fi','Private local network\nFirewall allows TCP 5080','#FFFFFF'),(1030,105,360,130,'Office browser','http://MAIN-PC-IP:5080','#FFFFFF'),(1030,285,360,130,'Mobile / laptop','Same database and live view','#FFFFFF')],[(390,260,560,260),(860,260,1030,170),(860,260,1030,350)])
    diagram('diagram-commissioning.png','Commissioning sequence',[
      (55,150,260,220,'1. Shop','Customer, location\nand invoice email','#FFFFFF'),(350,150,260,220,'2. Controller','RTU COM port or\nTCP gateway address','#FFFFFF'),(645,150,260,220,'3. Meter','Shop + source +\nregister/scale','#FFFFFF'),(940,150,220,220,'4. Test','Verify live value\nand status','#DFF3F1'),(1195,150,250,220,'5. Activate','Confirm commissioned\nthen monitor','#FFFFFF')],[(315,260,350,260),(610,260,645,260),(905,260,940,260),(1160,260,1195,260)])

def cover(d):
    sec=d.sections[0]; sec.header.is_linked_to_previous=False
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; picture(p.add_run(),IMG/'watch-dog-logo.png',2.0,'Watch Dog Energy Management logo')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('WATCH DOG'); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(TEAL)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('ENERGY MANAGEMENT'); r.bold=True; r.font.size=Pt(30); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Customer User Guide'); r.font.size=Pt(22); r.font.color.rgb=RGBColor.from_string(TEAL)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Watch Every Watt').italic=True
    d.add_paragraph('\n')
    simple_table(d,['Document','Details'],[['Software release','2.2.1'],['Guide edition','September 2026'],['Audience','Mall operators, billing staff and administrators'],['Prepared by','MicroBrain by Fadi Assi']],[1.8,4.8])
    note(d,'Purpose','Operate, monitor and bill reliably from one Watch Dog installation. Screens show fictional demonstration data; site values will differ.')
    d.add_paragraph('\n')
    p=d.add_paragraph('Customer copy  •  Keep near the main Watch Dog PC'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def toc(d):
    page(d); label(d,'Document map'); title(d,'Contents','A practical route from commissioning through monitoring, billing and support.')
    rows=[('1','System at a glance'),('2','Start, sign in and navigate'),('3','Dashboard'),('4','Grid and Generator pages'),('5','Meter details and consumption'),('6','Shops, controllers and meters'),('7','Optional Water monitoring'),('8','Tariffs and monthly publication'),('9','Reading history and data export'),('10','Published invoices'),('11','Invoice email (optional)'),('12','Users, access and audit'),('13','System features, backup and restore'),('14','LAN browser access'),('15','Daily operating checklist'),('16','Troubleshooting'),('17','FAQ and handover')]
    simple_table(d,['Section','Topic'],rows,[.9,5.7])
    note(d,'Terminology','The current interface says Grid for utility electricity. “UMEME” was the name used by earlier releases.')

def section_intro(d,num,name,sub):
    page(d); label(d,f'Section {num}'); title(d,name,sub)

def build_manual():
    make_diagrams(); d=base_doc(); cover(d); toc(d)
    section_intro(d,1,'System at a glance','What Watch Dog does—and what remains in the field meter.')
    screenshot(d,'diagram-system.png','Figure 1 — End-to-end operating model',7.0)
    d.add_heading('What the system does',1)
    bullet(d,'Reads cumulative utility values supplied by commissioned meters through Modbus RTU or Modbus TCP gateways.')
    bullet(d,'Stores time-stamped readings in one local database, shows status and consumption trends, and calculates billing-period consumption.')
    bullet(d,'Applies the tariff saved for each meter, publishes monthly invoices, and keeps a searchable invoice archive.')
    bullet(d,'Optionally emails the published PDF invoice through the customer’s own SMTP service.')
    note(d,'Important','Watch Dog does not replace a certified energy meter and does not calculate the physical energy quantity. Billing is based on the reliable cumulative reading received from the meter.')

    section_intro(d,2,'Start, sign in and navigate','The same application is used on the main PC and authorized browser clients.')
    screenshot(d,'01-login.png','Figure 2 — Administrator sign-in from the current page',7.0)
    d.add_heading('Start and sign in',2)
    numbered(d,'On the main PC, start Watch Dog Energy Management using the installed shortcut or configured startup method.')
    numbered(d,'Wait for the browser to open at http://localhost:5080.')
    numbered(d,'Use the Administrator level link when configuration is required, then enter the assigned username and password.')
    numbered(d,'Use Main for operational pages and Settings for administrator pages. Log out when finished.')
    note(d,'Access levels','Administrator: full setup. Billing Manager: tariffs, invoices and payments. Operator: monitoring and routine operations. Viewer: read-only monitoring. Visible pages depend on the signed-in role.')

    section_intro(d,3,'Dashboard','One operational view of electricity meters, invoices and system condition.')
    screenshot(d,'02-dashboard.png','Figure 3 — Dashboard with fictional demonstration data',7.0)
    simple_table(d,['Area','How to use it'],[
      ['Summary cards','Check shops, active meters, meters online, alarms and outstanding balances.'],['Tariff banner','Confirms the range and count of active meter-specific electricity tariffs.'],['Live meter status','Shows Grid/Generator source, MDB panel, accumulated reading, communication state and last-reading time.'],['Invoice register','Opens recent invoices and their PDF/detail page.'],['System condition','Use status totals to decide whether to investigate communication or billing.']],[1.65,5.0])
    note(d,'Healthy reading','“Connected” plus a recent Last reading time is the strongest quick check. A value that never changes may still be a valid low-load period; compare timestamps and the physical meter before changing settings.')

    section_intro(d,4,'Grid and Generator pages','Source-specific electricity monitoring with the same search and meter workflow.')
    screenshot(d,'03-grid.png','Figure 4 — Grid page: only utility-supply meters',7.0)
    screenshot(d,'04-generator.png','Figure 5 — Generator page: only generator-supply meters',7.0)
    numbered(d,'Open Main → Grid or Main → Generator.')
    numbered(d,'Search by meter ID/number, meter name, customer or tenant, shop, room or location.')
    numbered(d,'Review accumulated kWh, controller, communication and latest timestamp.')
    numbered(d,'Select View meter for detailed Per Day, Per Month and Per Year consumption.')
    note(d,'Automatic placement','A meter appears on the correct page according to the Power source selected when its meter record is created or edited.')

    section_intro(d,5,'Meter details and consumption','Turn cumulative meter readings into understandable hourly, daily and monthly usage.')
    screenshot(d,'06-meter-details.png','Figure 6 — Per Day, Per Month and Per Year consumption charts',5.35)
    d.add_heading('How to read the charts',2)
    simple_table(d,['Chart','Each bar represents','Typical use'],[['Per Day','Consumption during one hour','Spot daily peaks or missing intervals.'],['Per Month','Consumption during one day','Compare days in the current month.'],['Per Year','Consumption during one month','Compare seasonal/monthly patterns.']],[1.25,2.25,3.15])
    bullet(d,'Move the pointer over a bar to see its period, consumption and cost when a valid tariff is available.')
    bullet(d,'Totals show consumption differences between readings—not the accumulated register multiplied by tariff.')
    note(d,'Data quality','A chart needs at least two usable readings around a period. Missing bars usually mean no suitable readings, not zero consumption.')

    section_intro(d,6,'Shops, controllers and meters','Administrator commissioning sequence for a new customer meter.')
    screenshot(d,'diagram-commissioning.png','Figure 7 — Recommended order prevents missing assignments',7.0)
    screenshot(d,'12-shop-setup.png','Figure 8 — Shop/customer records can be created and edited',7.0)
    d.add_heading('Step 1 — Create the shop',2)
    numbered(d,'Open Settings → Shops and choose Add shop.')
    numbered(d,'Enter shop number, name, tenant/customer, contact, invoice email and location. Save.')
    d.add_heading('Step 2 — Create and test the controller',2)
    numbered(d,'Open Settings → Controller setup and choose Add controller.')
    numbered(d,'Select Modbus RTU for a Windows COM port, or Modbus TCP for an IP gateway/controller.')
    numbered(d,'Enter the field settings exactly as commissioned, save, then use Edit / test.')
    screenshot(d,'13-controller-setup.png','Figure 9 — Controller list and communication status',7.0)
    note(d,'RTU versus TCP','RTU uses a Windows COM port for the USB/RS-485 adapter; the DDC’s physical port name does not have to equal the Windows COM number. TCP uses the controller/gateway IP address and port, normally 502.')

    page(d); label(d,'Section 6 continued'); title(d,'Create and commission a meter','Every meter belongs to a shop, controller/gateway and utility/source.')
    d.add_heading('Step 3 — Add the meter',2)
    numbered(d,'From Settings → Controller setup, open the controller’s meter channels or add-meter action.')
    numbered(d,'Enter the meter name and actual serial number, then select the shop/customer.')
    numbered(d,'For electricity choose Grid or Generator. For water choose Water when the Water feature is enabled.')
    numbered(d,'Enter the register address, data type, word order and scale exactly as specified by the commissioned meter or gateway.')
    numbered(d,'Save, test communication, confirm the displayed value against the physical meter, then activate/commission it.')
    note(d,'Scale example','A scale of 0.01 means raw register value 123456 is displayed as 1,234.56 units. Never guess the scale—use the meter/gateway register documentation and verify against the meter display.')
    note(d,'Safety','Do not mark a meter commissioned until the identifier, shop assignment, source, unit, live value and timestamp have all been checked.')

    section_intro(d,7,'Optional Water monitoring','Water pages use the same workflow with blue styling and m³ units.')
    screenshot(d,'05-water.png','Figure 10 — Optional Water page with a Modbus TCP gateway',7.0)
    bullet(d,'Enable Water in Settings → System features only when water meters are part of the installation.')
    bullet(d,'An M-Bus meter may connect to an M-Bus-to-Modbus TCP gateway; Watch Dog reads the mapped Modbus value directly.')
    bullet(d,'Create the gateway as Modbus TCP, then create each water meter with its gateway unit/slave ID, register, data type, word order and scale.')
    note(d,'Commissioning check','Compare Watch Dog m³ with the physical water meter. Confirm that the gateway exposes a cumulative volume register and not an instantaneous flow register.')

    section_intro(d,8,'Tariffs and monthly publication','One tariff and currency per meter; one selected publication day per meter schedule.')
    screenshot(d,'07-tariff-schedule.png','Figure 11 — Tariff and monthly invoice publication settings',7.0)
    numbered(d,'Open Main → Operations and expand Tariff and monthly invoice publication.')
    numbered(d,'Select the meter and enter its price per kWh (or m³ for water) and ISO currency.')
    numbered(d,'Select an invoice publication day from 1 to 7 and enter the payment due interval in days.')
    numbered(d,'Choose Save all settings. The tariff starts immediately for future consumption calculations.')
    bullet(d,'On the scheduled day, Watch Dog publishes the previous completed calendar month when it is running.')
    bullet(d,'Each meter can have a different tariff, currency, publication day and due interval.')
    note(d,'Example','A meter scheduled for day 6 with 10 due days publishes August consumption on 6 September and sets the due date 10 days later. It does not invoice accumulated lifetime kWh.')

    section_intro(d,9,'Reading history and data export','Verify collected evidence before investigating billing or communication.')
    screenshot(d,'08-reading-history.png','Figure 12 — Reading history can be folded to keep Operations compact',4.55)
    bullet(d,'Filter readings by date/meter where controls are available, and use the latest timestamps to confirm collection continuity.')
    bullet(d,'Export readings to CSV for external review or reconciliation.')
    bullet(d,'A valid reading includes meter identity, timestamp, value and quality/state. Invalid or repeated data should be investigated before billing.')
    note(d,'Purpose','Reading history is the evidence behind consumption differences and invoices. The Audit trail records user/configuration events; it is not a substitute for meter readings.')

    section_intro(d,10,'Published invoices','Search, review, download, print and record payment against published invoices.')
    screenshot(d,'10-invoice-archive.png','Figure 13 — Searchable invoice archive',7.0)
    numbered(d,'Open Main → Published invoices.')
    numbered(d,'Search by day, month, year, shop/tenant, meter or invoice number.')
    numbered(d,'Select View / PDF, confirm the billing period, readings, tariff, charges, total and balance.')
    screenshot(d,'11-invoice-details.png','Figure 14 — Invoice detail and PDF/email actions',6.6)
    bullet(d,'Download PDF creates the customer document. Print preview uses the browser print workflow.')
    bullet(d,'Authorized staff can record payments from Operations; status changes to partially paid or paid according to the remaining balance.')
    note(d,'Before delivery','Confirm customer name/email, meter serial number, period, opening and closing readings, tariff/currency, due date and total.')

    section_intro(d,11,'Invoice email (optional)','Use the customer’s SMTP provider; the feature is disabled by default.')
    screenshot(d,'15-invoice-email.png','Figure 15 — SMTP settings and global automatic-delivery switch',7.0)
    numbered(d,'Open Settings → Invoice email as Administrator.')
    numbered(d,'Enter SMTP server, port, SSL/TLS mode, username, app password, sender name/email and a test recipient.')
    numbered(d,'Write the standard customer message. Watch Dog adds the greeting, invoice details and PDF attachment.')
    numbered(d,'Choose Save and send test email. Confirm receipt before enabling automatic delivery.')
    numbered(d,'Enable PDF invoice email. Also enable Automatically email every newly published invoice only when the customer wants automatic delivery.')
    note(d,'Password behavior','The SMTP/app password is encrypted for this Windows installation and is not displayed again. Leave the password box blank when changing other settings; enter a new value only to replace it.')
    note(d,'If test email fails','Check internet access, SMTP host/port, SSL/TLS, username, app-password validity, sender permission, antivirus/firewall and the provider’s sent/spam logs.')

    section_intro(d,12,'Users, access and audit','Give every person the minimum access needed and keep changes attributable.')
    screenshot(d,'14-users-access.png','Figure 16 — Users & Access is managed by an Administrator',6.8)
    numbered(d,'Open Settings → Users & Access, enter username, optional information email, access level and a temporary password.')
    numbered(d,'Create the user and give credentials securely. Test the account before operational handover.')
    numbered(d,'Delete accounts that are no longer required only after confirming the exact user.')
    simple_table(d,['Role','Use'],[['Administrator','System setup, communication, users, SMTP, backup and all operations.'],['Billing Manager','Tariffs, invoices and payment operations.'],['Operator','Daily monitoring and permitted operating tasks.'],['Viewer','Read-only monitoring.']],[1.6,5.0])
    screenshot(d,'17-audit-trail.png','Figure 17 — Audit trail keeps administrative and billing changes separate from daily screens',7.0)

    section_intro(d,13,'System features, backup and restore','Optional modules can be hidden; configuration and billing data must be protected.')
    screenshot(d,'16-system-features.png','Figure 18 — Enable Water only when required',7.0)
    screenshot(d,'09-backup-restore.png','Figure 19 — Backup and restore controls in Operations',7.0)
    numbered(d,'Create a backup after successful commissioning and before major configuration or software changes.')
    numbered(d,'Copy approved backups to a separate protected drive or server according to the customer’s policy.')
    numbered(d,'Use Restore only as Administrator, select the correct backup, and verify shops, meters, readings and invoices afterwards.')
    note(d,'Restore warning','Restore changes system data. Stop routine changes, identify the exact backup and ensure a separate current backup exists before proceeding.')

    section_intro(d,14,'LAN browser access','Other computers monitor the exact same Watch Dog database through the main PC.')
    screenshot(d,'diagram-lan.png','Figure 20 — Browser clients do not poll meters or run another backend',7.0)
    numbered(d,'Give the main Watch Dog PC a stable LAN address, for example 192.168.16.201, using a DHCP reservation or careful static configuration.')
    numbered(d,'On the main PC, confirm http://localhost:5080 opens and Windows Firewall has an inbound TCP rule for port 5080 on the Private profile.')
    numbered(d,'From a second device on the same LAN/Wi-Fi, open http://MAIN-PC-IP:5080—for example http://192.168.16.201:5080.')
    numbered(d,'If it fails, verify both devices are on the same subnet, the network profile is Private, client isolation is off, and ping/routing/firewall policy allows access.')
    note(d,'Architecture','Only the main PC communicates with the DDC/PLC/gateways. Multiple browsers share the same web server and database; they do not create additional polling services.')

    section_intro(d,15,'Daily operating checklist','A short routine for the operator at the start and end of each shift.')
    d.add_heading('Start of day',2)
    for x in ['Main PC is powered, awake and connected to the LAN and field communication adapters.','Watch Dog opens at localhost:5080.','Meters Online and Meter Alarms are reviewed.','Grid, Generator and enabled Water pages show recent timestamps.','Any offline or stale meter is recorded and investigated.']: bullet(d,x)
    d.add_heading('Billing-day check',2)
    for x in ['Published invoice archive contains the expected period and shops.','Invoices have correct opening/closing readings, tariff, currency and due date.','Automatic email results are checked when enabled; failed deliveries are handled manually.','Payments entered by staff reconcile with the remaining balances.']: bullet(d,x)
    d.add_heading('End of day / weekly',2)
    for x in ['Communication alarms are resolved or handed over.','Audit trail is reviewed for unexpected configuration changes.','A recent verified backup exists outside the main PC.']: bullet(d,x)

    section_intro(d,16,'Troubleshooting','Use evidence first; change configuration only when the fault is identified.')
    simple_table(d,['Symptom','Check','Action'],[
      ['Not commissioned','Controller/meter enabled, communication test and activation','Test settings, compare the live value, then confirm commissioning.'],
      ['No communication','Windows COM/IP, wiring, polarity, baud/parity/slave or TCP route/port','Match field settings exactly; ensure no second program owns the COM port.'],
      ['Connected but value does not change','Timestamp, meter display and actual load','Wait for load/change; verify the correct cumulative register and scale.'],
      ['No invoice','Schedule day, app uptime, readings around month boundary, tariff','Keep main PC awake/running; verify two suitable readings and active tariff.'],
      ['Email test fails','SMTP/app password, TLS, port, internet and provider policy','Re-enter new app password, save-and-test, check provider logs.'],
      ['Second PC cannot open','Main PC IP, port 5080, Private firewall rule, Wi-Fi isolation','Test localhost first, then http://MAIN-PC-IP:5080.'],
      ['Charts empty','Insufficient historical reading points','Confirm timestamps and allow readings to accumulate.'],
      ['Water page missing','Feature disabled','Administrator enables Water in System features.']],[1.4,2.5,2.75])
    note(d,'Escalate with evidence','Provide the exact meter/controller, timestamp, visible status, recent change, and screenshot. Never send passwords or unredacted customer secrets.')

    section_intro(d,17,'FAQ and handover','Answers to common customer questions.')
    faqs=[
      ('Does each meter have its own tariff?','Yes. Select the meter, price and currency in Operations. Tariff history is retained through the audit trail.'),
      ('Does the tariff date create the invoice?','The current workflow uses a monthly publication day from 1–7. The tariff becomes active immediately; the schedule publishes the previous completed month.'),
      ('Why is accumulated kWh larger than the monthly bill usage?','Accumulated reading is the lifetime register. Invoice consumption is closing minus opening reading for the billing period.'),
      ('Can Grid and Generator meters use different prices?','Yes. Tariff and currency are saved per meter.'),
      ('Can customers view from a phone?','Yes, on the same LAN/Wi-Fi using the main PC IP and port 5080, subject to login and network policy.'),
      ('Does every browser communicate with the PLC?','No. One background service on the main PC polls meters; browsers only view the same application.'),
      ('What happens if the PC sleeps?','Polling and scheduled work cannot run while Windows is asleep. Keep the main PC powered and configure an appropriate no-sleep policy.'),
      ('Can an invoice be emailed automatically?','Yes, when SMTP is correctly tested and both PDF invoice email and automatic delivery are enabled.'),
      ('What should be backed up?','Use Watch Dog’s backup function; retain protected copies away from the main PC and test restore procedures.'),
    ]
    for q,a in faqs: d.add_heading(q,2); d.add_paragraph(a)
    d.add_heading('Customer handover record',1)
    simple_table(d,['Item','Confirmed by / date'],[['Main PC and LAN URL',''],['Administrator account issued securely',''],['Controllers and meters tested',''],['Tariffs and schedules approved',''],['SMTP test received (if enabled)',''],['Backup created and copied',''],['Operator training completed','']],[3.8,2.8])
    note(d,'Support boundary','Field wiring, meter programming, Windows/network policy and SMTP provider accounts remain site/customer responsibilities unless included in a separate support agreement.')
    return d

def quick_guide():
    d=base_doc(); p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; picture(p.add_run(),IMG/'watch-dog-logo.png',.8,'Watch Dog Energy Management logo'); label(d,'Watch Dog Energy Management'); title(d,'Quick User Guide','Daily monitoring, billing and first-response troubleshooting • Version 2.2.1')
    simple_table(d,['Every day','What to confirm'],[['Open','Main PC: http://localhost:5080  •  LAN client: http://MAIN-PC-IP:5080'],['Dashboard','Meters online, alarms, recent timestamps, correct Grid/Generator source'],['Meter','View meter → Per Day / Per Month / Per Year; hover bars for usage/cost'],['Invoice','Published invoices → search → View / PDF → verify → deliver'],['Backup','Confirm a recent protected backup exists']],[1.25,5.4])
    d.add_heading('Add a new meter — correct order',1)
    for s in ['1. Settings → Shops: create the customer/shop.','2. Settings → Controller setup: add RTU COM or TCP gateway and test.','3. Add meter: identity + shop + Grid/Generator/Water + register/data type/word order/scale.','4. Test: compare Watch Dog with the physical meter, then activate/commission.','5. Operations: select meter, tariff/currency, publication day 1–7 and due days; Save all settings.']: bullet(d,s)
    note(d,'Never guess','Use the commissioned meter/gateway register information. A wrong register, word order or scale can produce a believable but incorrect bill.', 'FFF3D7')
    d.add_heading('No invoice on the selected day?',1)
    bullet(d,'Confirm Watch Dog was running and the PC was awake.')
    bullet(d,'Confirm the tariff and publication schedule were saved for that meter.')
    bullet(d,'Confirm readings exist around the start and end of the previous completed month.')
    bullet(d,'Check Published invoices and the Audit trail; then record evidence before changing settings.')
    d.add_heading('No communication?',1)
    bullet(d,'RTU: Windows COM port, RS-485 polarity/wiring, baud, parity, data bits, stop bits and slave address.')
    bullet(d,'TCP: controller/gateway IP, TCP port, route, unit/slave address and one reachable LAN.')
    bullet(d,'Check the last-reading timestamp and compare the physical meter value.')
    d.add_heading('Escalation note',1); d.add_paragraph('Record meter/controller name, exact time, status, recent change and screenshot. Do not send passwords or secrets.')
    return d

if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    manual=build_manual(); manual.save(OUT/'Watch_Dog_EM_Customer_User_Guide.docx')
    q=quick_guide(); q.save(OUT/'Watch_Dog_EM_Quick_User_Guide.docx')
    print('DOCX_CREATED')
